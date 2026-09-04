import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
from openpyxl.worksheet.properties import WorksheetProperties, PageSetupProperties
from copy import copy
import datetime

# ---------------------------------------------------------------------------
# PALETTE / FONTS
# ---------------------------------------------------------------------------
TEAL      = "2F5D62"   # primary dark teal
TEAL_LT   = "DCEAE8"   # light teal tint (section fills)
GOLD      = "D9A441"   # accent gold
GOLD_LT   = "F7E9CC"
BLUSH     = "E8B4B8"
BLUSH_LT  = "FBEAEB"
SAGE      = "8FA998"
SAGE_LT   = "E7EEE9"
CREAM     = "FBF7F0"
CHARCOAL  = "3A3A3A"
WHITE     = "FFFFFF"
INPUT_BLUE = "1F4E78"   # font color for user-input cells
RED       = "C0504D"
GREEN     = "4F7942"

FONT_NAME = "Arial"

def F(size=11, bold=False, italic=False, color=CHARCOAL, name=FONT_NAME):
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)

def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)

THIN = Side(style="thin", color="C9C9C9")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=Side(style="medium", color=TEAL))

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

CURRENCY_FMT = '$#,##0.00;[RED]-$#,##0.00'
CURRENCY_FMT0 = '$#,##0;[RED]-$#,##0'
PCT_FMT = '0.0%'
DATE_FMT = 'mm/dd/yyyy'

MONTHS = ["January","February","March","April","May","June",
          "July","August","September","October","November","December"]
MONTH_ABBR = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]

INCOME_CATS = ["Salary","Freelance / Side Hustle","Bonus","Investment Income","Gifts Received","Other Income"]
FIXED_CATS = ["Rent / Mortgage","Utilities","Insurance","Phone Bill","Internet","Subscriptions","Car Payment","Loan Payments"]
VARIABLE_CATS = ["Groceries","Dining Out","Transportation / Gas","Entertainment","Shopping",
                  "Personal Care","Health / Medical","Pet Care","Childcare","Gifts & Donations","Travel","Miscellaneous"]
SAVINGS_CATS = ["Emergency Fund","Retirement Savings","Investment Contributions","Debt Payoff","Vacation Fund","Other Savings"]

ALL_SECTIONS = [
    ("INCOME", INCOME_CATS, TEAL, TEAL_LT, "Income"),
    ("FIXED EXPENSES", FIXED_CATS, GOLD, GOLD_LT, "Fixed Expense"),
    ("VARIABLE EXPENSES", VARIABLE_CATS, BLUSH, BLUSH_LT, "Variable Expense"),
    ("SAVINGS & DEBT", SAVINGS_CATS, SAGE, SAGE_LT, "Savings & Debt"),
]

wb = Workbook()
wb.remove(wb.active)

def style_sheet_base(ws, gridlines=False, tab_color=None, freeze=None, zoom=100):
    ws.sheet_view.showGridLines = gridlines
    ws.sheet_view.zoomScale = zoom
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    if freeze:
        ws.freeze_panes = freeze

def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def title_block(ws, text, subtitle=None, span="A1:D1", color=TEAL, fontsize=20, height=34):
    ws.merge_cells(span)
    c = ws[span.split(":")[0]]
    c.value = text
    c.font = F(size=fontsize, bold=True, color=WHITE)
    c.fill = fill(color)
    c.alignment = LEFT
    ws.row_dimensions[int(span.split(":")[0][1:])].height = height
    # fill entire row across merge with color already via merge; ensure right padding cell same fill
    if subtitle:
        row2 = int(span.split(":")[0][1:]) + 1
        span2 = span.split(":")[0][0] + str(row2) + ":" + span.split(":")[1][0] + str(row2)
        ws.merge_cells(span2)
        c2 = ws[span2.split(":")[0]]
        c2.value = subtitle
        c2.font = F(size=10, italic=True, color=CHARCOAL)
        c2.alignment = LEFT

# ---------------------------------------------------------------------------
# WELCOME / COVER SHEET
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Welcome")
style_sheet_base(ws, tab_color=TEAL)
set_col_widths(ws, {"A": 3, "B": 30, "C": 55, "D": 3})
ws.sheet_view.showGridLines = False

for r in range(1, 40):
    for c in range(1, 5):
        ws.cell(row=r, column=c).fill = fill(CREAM)

ws.merge_cells("B2:C3")
ws["B2"] = "THE ULTIMATE\nBUDGET PLANNER"
ws["B2"].font = F(size=26, bold=True, color=TEAL)
ws["B2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 30
ws.row_dimensions[3].height = 30

ws.merge_cells("B4:C4")
ws["B4"] = "A complete system for planning, tracking & growing your money"
ws["B4"].font = F(size=12, italic=True, color=GOLD)

ws.merge_cells("B6:C6")
ws["B6"] = "YEAR"
ws["B6"].font = F(size=10, bold=True, color=WHITE)
ws["B6"].fill = fill(TEAL)
ws["B6"].alignment = CENTER
ws.merge_cells("B7:C7")
ws["B7"] = 2026
ws["B7"].font = F(size=16, bold=True, color=INPUT_BLUE)
ws["B7"].alignment = CENTER
ws["B7"].fill = fill(GOLD_LT)
ws.row_dimensions[7].height = 22
wb_year_cell = "Welcome!$B$7"

def section_header(row, text):
    ws.merge_cells(f"B{row}:C{row}")
    cell = ws[f"B{row}"]
    cell.value = text
    cell.font = F(size=13, bold=True, color=WHITE)
    cell.fill = fill(GOLD)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 22

section_header(9, "WHAT'S INSIDE")
items = [
    ("Dashboard", "Your at-a-glance yearly overview with automatic charts."),
    ("Jan - Dec Tabs", "A full monthly budget: plan it, then track what actually happened."),
    ("Transactions", "Log every purchase once — everything else updates itself."),
    ("Annual Overview", "All 12 months side-by-side, so trends jump right out."),
    ("Debt Payoff Tracker", "See exactly when each debt will hit zero."),
    ("Savings Goals", "Turn 'someday' goals into a monthly savings number."),
]
r = 10
for name, desc in items:
    ws[f"B{r}"] = name
    ws[f"B{r}"].font = F(size=11, bold=True, color=TEAL)
    ws[f"C{r}"] = desc
    ws[f"C{r}"].font = F(size=10.5, color=CHARCOAL)
    ws[f"C{r}"].alignment = LEFT_WRAP
    ws.row_dimensions[r].height = 28
    r += 1

section_header(r+1, "HOW TO USE THIS PLANNER")
r += 2
steps = [
    "1.  Start on the Transactions tab and log your income & expenses as they happen (or all at once for the month).",
    "2.  Open the matching month tab (Jan, Feb, ...) and fill in the yellow BUDGETED column with your plan for each category.",
    "3.  The ACTUAL and VARIANCE columns fill in automatically from your Transactions log — no manual math required.",
    "4.  Check the Dashboard any time for a live picture of the whole year.",
    "5.  Use the Debt Payoff and Savings Goals tabs to plan beyond the month-to-month.",
]
for s in steps:
    ws.merge_cells(f"B{r}:C{r}")
    ws[f"B{r}"] = s
    ws[f"B{r}"].font = F(size=10.5, color=CHARCOAL)
    ws[f"B{r}"].alignment = LEFT_WRAP
    ws.row_dimensions[r].height = 26
    r += 1

section_header(r+1, "COLOR KEY")
r += 2
legend = [
    (GOLD_LT, "Yellow / gold cells = type your own numbers here"),
    (WHITE,   "White cells with a formula = calculated automatically, don't overwrite"),
    (TEAL_LT, "Teal band = Income section"),
    (BLUSH_LT,"Blush band = Variable Expenses section"),
]
for color, text in legend:
    ws[f"B{r}"].fill = fill(color)
    ws[f"B{r}"].border = BORDER_ALL
    ws.merge_cells(f"C{r}:C{r}")
    ws[f"C{r}"] = text
    ws[f"C{r}"].font = F(size=10, color=CHARCOAL)
    ws.row_dimensions[r].height = 18
    r += 1

r += 2
ws.merge_cells(f"B{r}:C{r}")
ws[f"B{r}"] = "Tip: right-click any sheet tab to duplicate it if you'd like extra copies (e.g. per bank account)."
ws[f"B{r}"].font = F(size=9.5, italic=True, color=SAGE)
ws[f"B{r}"].alignment = LEFT_WRAP

# ---------------------------------------------------------------------------
# CATEGORIES (lookup table powering dropdowns + auto-type lookups)
# ---------------------------------------------------------------------------
cat = wb.create_sheet("Categories")
style_sheet_base(cat, tab_color=CHARCOAL)
set_col_widths(cat, {"A": 28, "B": 20, "C": 3})
cat["A1"] = "CATEGORY"
cat["B1"] = "TYPE"
for col in ("A1", "B1"):
    cat[col].font = F(bold=True, color=WHITE)
    cat[col].fill = fill(TEAL)
    cat[col].alignment = CENTER

row = 2
CATEGORY_ROWS = {}   # category name -> row number on Categories sheet
for _, cats, _, _, typelabel in ALL_SECTIONS:
    for cname in cats:
        cat.cell(row=row, column=1, value=cname).font = F()
        cat.cell(row=row, column=2, value=typelabel).font = F(color=CHARCOAL)
        CATEGORY_ROWS[cname] = row
        row += 1
LAST_CAT_ROW = row - 1
for r in range(2, LAST_CAT_ROW + 1):
    cat.cell(row=r, column=1).border = BORDER_ALL
    cat.cell(row=r, column=2).border = BORDER_ALL

cat.sheet_state = "visible"

ALL_CATS_RANGE = f"Categories!$A$2:$A${LAST_CAT_ROW}"
TYPE_LOOKUP_RANGE_CAT = f"Categories!$A$2:$A${LAST_CAT_ROW}"
TYPE_LOOKUP_RANGE_TYPE = f"Categories!$B$2:$B${LAST_CAT_ROW}"

print("welcome + categories ok, last cat row", LAST_CAT_ROW)

# ---------------------------------------------------------------------------
# TRANSACTIONS LOG
# ---------------------------------------------------------------------------
tx = wb.create_sheet("Transactions")
style_sheet_base(tx, tab_color=GOLD, freeze="A4")
set_col_widths(tx, {"A": 12, "B": 30, "C": 24, "D": 16, "E": 14, "F": 9, "G": 18, "H": 26})

title_block(tx, "TRANSACTIONS LOG", "Log every purchase and paycheck here — the whole planner updates itself from this sheet.",
            span="A1:H1", color=GOLD, fontsize=16, height=30)

HEADERS = ["Date","Description","Category","Type","Amount","Month#","Payment Method","Notes"]
for i, h in enumerate(HEADERS, start=1):
    c = tx.cell(row=3, column=i, value=h)
    c.font = F(bold=True, color=WHITE)
    c.fill = fill(TEAL)
    c.alignment = CENTER
    c.border = BORDER_ALL
tx.row_dimensions[3].height = 20

TX_FIRST_DATA_ROW = 4
TX_LAST_ROW = 400  # generous room for the year

# Sample transactions (realistic January example) - clearly an example, user deletes/edits
sample = [
    (datetime.date(2026,1,1), "January Paycheck", "Salary", "Direct Deposit"),
    (datetime.date(2026,1,3), "Rent Payment", "Rent / Mortgage", "Bank Transfer"),
    (datetime.date(2026,1,4), "Electric & Gas Bill", "Utilities", "Debit Card"),
    (datetime.date(2026,1,5), "Trader Joe's", "Groceries", "Credit Card"),
    (datetime.date(2026,1,7), "Netflix + Spotify", "Subscriptions", "Credit Card"),
    (datetime.date(2026,1,10), "Gas Station", "Transportation / Gas", "Debit Card"),
    (datetime.date(2026,1,12), "Dinner with friends", "Dining Out", "Credit Card"),
    (datetime.date(2026,1,15), "January Paycheck", "Salary", "Direct Deposit"),
    (datetime.date(2026,1,15), "Transfer to Savings", "Emergency Fund", "Bank Transfer"),
    (datetime.date(2026,1,18), "Credit Card Payment", "Debt Payoff", "Bank Transfer"),
    (datetime.date(2026,1,22), "Freelance Project", "Freelance / Side Hustle", "PayPal"),
    (datetime.date(2026,1,27), "Weekly Grocery Run", "Groceries", "Credit Card"),
]
sample_amounts = [3200.00, 1450.00, 165.32, 98.47, 24.98, 45.00, 62.10, 3200.00, 300.00, 150.00, 450.00, 112.60]

r = TX_FIRST_DATA_ROW
for (d, desc, cname, method), amt in zip(sample, sample_amounts):
    tx.cell(row=r, column=1, value=d).number_format = DATE_FMT
    tx.cell(row=r, column=2, value=desc)
    tx.cell(row=r, column=3, value=cname)
    tx.cell(row=r, column=5, value=amt).number_format = CURRENCY_FMT
    tx.cell(row=r, column=7, value=method)
    r += 1

# Fill formulas + formatting down the full usable range
for rr in range(TX_FIRST_DATA_ROW, TX_LAST_ROW + 1):
    # Type auto-looked-up from Categories sheet based on chosen Category
    tx.cell(row=rr, column=4, value=f'=IFERROR(INDEX({TYPE_LOOKUP_RANGE_TYPE},MATCH(C{rr},{TYPE_LOOKUP_RANGE_CAT},0)),"")')
    # Month number helper, blank if no date
    tx.cell(row=rr, column=6, value=f'=IF(A{rr}="","",MONTH(A{rr}))')
    for col in range(1, 9):
        cell = tx.cell(row=rr, column=col)
        cell.font = F(size=10.5, color=INPUT_BLUE if col in (1,2,3,5,7,8) else CHARCOAL)
        cell.border = BORDER_ALL
        if col == 1:
            cell.number_format = DATE_FMT
            cell.alignment = CENTER
        if col == 5:
            cell.number_format = CURRENCY_FMT
        if col in (4,6):
            cell.alignment = CENTER
            cell.font = F(size=10.5, color=CHARCOAL, italic=True)

# alternating row shading for readability
for rr in range(TX_FIRST_DATA_ROW, TX_LAST_ROW + 1):
    if (rr - TX_FIRST_DATA_ROW) % 2 == 1:
        for col in range(1, 9):
            tx.cell(row=rr, column=col).fill = fill("F5F5F0")

# Data validation dropdowns
dv_cat = DataValidation(type="list", formula1=f"={ALL_CATS_RANGE}", allow_blank=True, showDropDown=False)
dv_cat.error = "Please choose a category from the dropdown list."
dv_cat.prompt = "Pick a category"
tx.add_data_validation(dv_cat)
dv_cat.add(f"C{TX_FIRST_DATA_ROW}:C{TX_LAST_ROW}")

dv_method = DataValidation(type="list", formula1='"Cash,Debit Card,Credit Card,Bank Transfer,Direct Deposit,PayPal,Check,Other"', allow_blank=True, showDropDown=False)
tx.add_data_validation(dv_method)
dv_method.add(f"G{TX_FIRST_DATA_ROW}:G{TX_LAST_ROW}")

TX_AMOUNT_RANGE = f"Transactions!$E${TX_FIRST_DATA_ROW}:$E${TX_LAST_ROW}"
TX_CATEGORY_RANGE = f"Transactions!$C${TX_FIRST_DATA_ROW}:$C${TX_LAST_ROW}"
TX_MONTHNUM_RANGE = f"Transactions!$F${TX_FIRST_DATA_ROW}:$F${TX_LAST_ROW}"

print("transactions ok")

# ---------------------------------------------------------------------------
# MONTHLY BUDGET SHEETS (Jan - Dec)
# ---------------------------------------------------------------------------
# Realistic example budget numbers for the fully-worked January sample sheet
EXAMPLE_BUDGET = {
    "Salary": 3200*2, "Freelance / Side Hustle": 400, "Bonus": 0, "Investment Income": 0,
    "Gifts Received": 0, "Other Income": 0,
    "Rent / Mortgage": 1450, "Utilities": 180, "Insurance": 120, "Phone Bill": 65,
    "Internet": 60, "Subscriptions": 35, "Car Payment": 320, "Loan Payments": 150,
    "Groceries": 450, "Dining Out": 150, "Transportation / Gas": 120, "Entertainment": 80,
    "Shopping": 100, "Personal Care": 50, "Health / Medical": 60, "Pet Care": 40,
    "Childcare": 0, "Gifts & Donations": 40, "Travel": 0, "Miscellaneous": 50,
    "Emergency Fund": 300, "Retirement Savings": 200, "Investment Contributions": 0,
    "Debt Payoff": 150, "Vacation Fund": 50, "Other Savings": 0,
}

MONTH_ROWS = {}   # month_abbr -> {"category_row": {cat: row}, "section_total_row": {section: row}, "summary": {label: row}}

def build_month_sheet(idx, month_name, month_abbr, is_example):
    ws = wb.create_sheet(month_abbr)
    style_sheet_base(ws, tab_color=TEAL if idx % 2 == 0 else GOLD, freeze="A5")
    set_col_widths(ws, {"A": 30, "B": 14, "C": 14, "D": 14, "E": 3})

    ws.merge_cells("A1:D1")
    ws["A1"] = f'="{month_name.upper()}"&" "&{wb_year_cell}&" BUDGET"'
    ws["A1"].font = F(size=17, bold=True, color=WHITE)
    ws["A1"].fill = fill(TEAL)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:D2")
    ws["A2"] = "Plan it in BUDGETED, then let ACTUAL fill in from your Transactions log."
    ws["A2"].font = F(size=10, italic=True, color=CHARCOAL)

    headers = ["CATEGORY", "BUDGETED", "ACTUAL", "VARIANCE"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = F(bold=True, color=WHITE)
        c.fill = fill(CHARCOAL)
        c.alignment = CENTER
        c.border = BORDER_ALL
    ws.row_dimensions[4].height = 20

    row = 5
    cat_rows = {}
    section_total_rows = {}
    section_ranges = {}  # section name -> (first_cat_row, last_cat_row)

    for section_name, cats, color, color_lt, _ in ALL_SECTIONS:
        ws.merge_cells(f"A{row}:D{row}")
        c = ws[f"A{row}"]
        c.value = section_name
        c.font = F(size=11, bold=True, color=WHITE)
        c.fill = fill(color)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row].height = 19
        row += 1
        first_row = row
        for cname in cats:
            ws.cell(row=row, column=1, value=cname).font = F(size=10.5)
            ws.cell(row=row, column=1).fill = fill(color_lt)
            ws.cell(row=row, column=1).border = BORDER_ALL

            budget_val = EXAMPLE_BUDGET[cname] if is_example else 0
            bcell = ws.cell(row=row, column=2, value=budget_val)
            bcell.number_format = CURRENCY_FMT0
            bcell.font = F(size=10.5, bold=True, color=INPUT_BLUE)
            bcell.fill = fill(GOLD_LT)
            bcell.border = BORDER_ALL
            bcell.alignment = RIGHT

            acell = ws.cell(row=row, column=3,
                value=f'=SUMIFS({TX_AMOUNT_RANGE},{TX_CATEGORY_RANGE},$A{row},{TX_MONTHNUM_RANGE},{idx+1})')
            acell.number_format = CURRENCY_FMT0
            acell.font = F(size=10.5, color=CHARCOAL)
            acell.fill = fill(color_lt)
            acell.border = BORDER_ALL
            acell.alignment = RIGHT

            vcell = ws.cell(row=row, column=4, value=f'=C{row}-B{row}')
            vcell.number_format = CURRENCY_FMT0
            vcell.font = F(size=10.5, color=CHARCOAL)
            vcell.fill = fill(color_lt)
            vcell.border = BORDER_ALL
            vcell.alignment = RIGHT

            cat_rows[cname] = row
            row += 1
        last_row = row - 1
        section_ranges[section_name] = (first_row, last_row)

        ws.cell(row=row, column=1, value=f"Total {section_name.title()}").font = F(size=10.5, bold=True, color=CHARCOAL)
        ws.cell(row=row, column=1).fill = fill(WHITE)
        for col, letter in ((2, "B"), (3, "C"), (4, "D")):
            cell = ws.cell(row=row, column=col, value=f"=SUM({letter}{first_row}:{letter}{last_row})")
            cell.number_format = CURRENCY_FMT0
            cell.font = F(size=10.5, bold=True, color=TEAL)
            cell.border = Border(top=Side(style="thin", color=CHARCOAL))
            cell.alignment = RIGHT
        ws.cell(row=row, column=1).border = Border(top=Side(style="thin", color=CHARCOAL))
        section_total_rows[section_name] = row
        row += 2  # blank spacer row

    # ---- SUMMARY BLOCK ----
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "MONTH SUMMARY"
    ws[f"A{row}"].font = F(size=11, bold=True, color=WHITE)
    ws[f"A{row}"].fill = fill(CHARCOAL)
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 19
    row += 1

    inc_total_row = section_total_rows["INCOME"]
    fixed_total_row = section_total_rows["FIXED EXPENSES"]
    var_total_row = section_total_rows["VARIABLE EXPENSES"]
    sav_total_row = section_total_rows["SAVINGS & DEBT"]

    summary_rows = {}
    def summary_line(label, budget_formula, actual_formula, bold=False, top_border=False):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = F(size=10.5, bold=bold)
        b = ws.cell(row=row, column=2, value=budget_formula)
        a = ws.cell(row=row, column=3, value=actual_formula)
        v = ws.cell(row=row, column=4, value=f"=C{row}-B{row}")
        for cell in (b, a, v):
            cell.number_format = CURRENCY_FMT0
            cell.font = F(size=10.5, bold=bold, color=TEAL if bold else CHARCOAL)
            cell.alignment = RIGHT
            if top_border:
                cell.border = Border(top=Side(style="thin", color=CHARCOAL))
        if top_border:
            ws.cell(row=row, column=1).border = Border(top=Side(style="thin", color=CHARCOAL))
        summary_rows[label] = row
        row += 1

    summary_line("Total Income", f"=B{inc_total_row}", f"=C{inc_total_row}")
    summary_line("Total Expenses (Fixed + Variable)",
                 f"=B{fixed_total_row}+B{var_total_row}", f"=C{fixed_total_row}+C{var_total_row}")
    summary_line("Total Savings & Debt Payments", f"=B{sav_total_row}", f"=C{sav_total_row}")
    total_income_row = summary_rows["Total Income"]
    total_exp_row = summary_rows["Total Expenses (Fixed + Variable)"]
    total_sav_row = summary_rows["Total Savings & Debt Payments"]
    summary_line("NET REMAINING (Income - Everything Else)",
                 f"=B{total_income_row}-B{total_exp_row}-B{total_sav_row}",
                 f"=C{total_income_row}-C{total_exp_row}-C{total_sav_row}",
                 bold=True, top_border=True)

    net_row = summary_rows["NET REMAINING (Income - Everything Else)"]
    for col_letter in ("B", "C", "D"):
        ws[f"{col_letter}{net_row}"].fill = fill(GOLD_LT)

    MONTH_ROWS[month_abbr] = {
        "cat_rows": cat_rows,
        "section_total_rows": section_total_rows,
        "summary_rows": summary_rows,
    }
    return ws

for idx, (mname, mabbr) in enumerate(zip(MONTHS, MONTH_ABBR)):
    build_month_sheet(idx, mname, mabbr, is_example=(idx == 0))

print("monthly sheets ok:", list(MONTH_ROWS.keys()))

# ---------------------------------------------------------------------------
# ANNUAL OVERVIEW
# ---------------------------------------------------------------------------
ao = wb.create_sheet("Annual Overview")
style_sheet_base(ao, tab_color=CHARCOAL, freeze="B5")

colw = {"A": 30}
for i in range(2, 14):
    colw[get_column_letter(i)] = 9.5
colw[get_column_letter(14)] = 12
colw[get_column_letter(15)] = 12
colw[get_column_letter(16)] = 12
set_col_widths(ao, colw)

ao.merge_cells("A1:P1")
ao["A1"] = "ANNUAL OVERVIEW — ACTUAL SPENDING BY MONTH"
ao["A1"].font = F(size=17, bold=True, color=WHITE)
ao["A1"].fill = fill(CHARCOAL)
ao["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ao.row_dimensions[1].height = 30

ao.merge_cells("A2:P2")
ao["A2"] = "Every number below is pulled automatically from the 12 monthly tabs — nothing to fill in here."
ao["A2"].font = F(size=10, italic=True, color=CHARCOAL)

header_labels = ["CATEGORY"] + MONTH_ABBR + ["TOTAL ACTUAL", "TOTAL BUDGETED", "VARIANCE"]
for i, h in enumerate(header_labels, start=1):
    c = ao.cell(row=4, column=i, value=h)
    c.font = F(size=9.5, bold=True, color=WHITE)
    c.fill = fill(TEAL)
    c.alignment = CENTER
    c.border = BORDER_ALL
ao.row_dimensions[4].height = 28

jan_layout = MONTH_ROWS["Jan"]
row = 5
ao_section_total_rows = {}
ao_cat_first_last = {}
AO_CAT_ROWS = {}   # category name -> row number on Annual Overview sheet

for section_name, cats, color, color_lt, _ in ALL_SECTIONS:
    ao.merge_cells(f"A{row}:P{row}")
    c = ao["A" + str(row)]
    c.value = section_name
    c.font = F(size=10.5, bold=True, color=WHITE)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ao.row_dimensions[row].height = 18
    row += 1
    first_row = row
    for cname in cats:
        month_row = jan_layout["cat_rows"][cname]
        ao.cell(row=row, column=1, value=cname).font = F(size=10)
        ao.cell(row=row, column=1).fill = fill(color_lt)
        ao.cell(row=row, column=1).border = BORDER_ALL
        for mi, mabbr in enumerate(MONTH_ABBR):
            cell = ao.cell(row=row, column=2 + mi, value=f"='{mabbr}'!C{month_row}")
            cell.number_format = CURRENCY_FMT0
            cell.font = F(size=9.5, color=CHARCOAL)
            cell.border = BORDER_ALL
            cell.alignment = RIGHT
        total_actual = ao.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})")
        total_budget = ao.cell(row=row, column=15,
            value="=" + "+".join(f"'{m}'!B{month_row}" for m in MONTH_ABBR))
        variance = ao.cell(row=row, column=16, value=f"=N{row}-O{row}")
        for cell in (total_actual, total_budget, variance):
            cell.number_format = CURRENCY_FMT0
            cell.font = F(size=9.5, bold=True, color=TEAL)
            cell.border = BORDER_ALL
            cell.alignment = RIGHT
        AO_CAT_ROWS[cname] = row
        row += 1
    last_row = row - 1
    ao_cat_first_last[section_name] = (first_row, last_row)

    ao.cell(row=row, column=1, value=f"Total {section_name.title()}").font = F(size=10, bold=True)
    for col in range(2, 17):
        letter = get_column_letter(col)
        cell = ao.cell(row=row, column=col, value=f"=SUM({letter}{first_row}:{letter}{last_row})")
        cell.number_format = CURRENCY_FMT0
        cell.font = F(size=9.5, bold=True, color=TEAL)
        cell.border = Border(top=Side(style="thin", color=CHARCOAL))
        cell.alignment = RIGHT
    ao.cell(row=row, column=1).border = Border(top=Side(style="thin", color=CHARCOAL))
    ao_section_total_rows[section_name] = row
    row += 2

# ---- Yearly summary rows (used by Dashboard charts) ----
ao.merge_cells(f"A{row}:P{row}")
ao[f"A{row}"] = "YEAR SUMMARY"
ao[f"A{row}"].font = F(size=10.5, bold=True, color=WHITE)
ao[f"A{row}"].fill = fill(CHARCOAL)
ao.row_dimensions[row].height = 18
row += 1

inc_row = ao_section_total_rows["INCOME"]
fixed_row = ao_section_total_rows["FIXED EXPENSES"]
var_row = ao_section_total_rows["VARIABLE EXPENSES"]
sav_row = ao_section_total_rows["SAVINGS & DEBT"]

AO_SUMMARY_ROWS = {}
def ao_summary(label, row_formula_per_month, bold=True):
    global row
    ao.cell(row=row, column=1, value=label).font = F(size=10, bold=bold)
    for mi in range(12):
        col = 2 + mi
        letter = get_column_letter(col)
        cell = ao.cell(row=row, column=col, value=row_formula_per_month(letter))
        cell.number_format = CURRENCY_FMT0
        cell.font = F(size=9.5, bold=bold, color=TEAL)
        cell.border = BORDER_ALL
        cell.alignment = RIGHT
    total_cell = ao.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})")
    total_cell.number_format = CURRENCY_FMT0
    total_cell.font = F(size=9.5, bold=True, color=TEAL)
    total_cell.border = BORDER_ALL
    AO_SUMMARY_ROWS[label] = row
    row += 1

ao_summary("Total Income", lambda l: f"={l}{inc_row}")
ao_summary("Total Expenses", lambda l: f"={l}{fixed_row}+{l}{var_row}")
ao_summary("Total Savings & Debt", lambda l: f"={l}{sav_row}")
net_label = "Net Cash Flow"
inc_r = AO_SUMMARY_ROWS["Total Income"]
exp_r = AO_SUMMARY_ROWS["Total Expenses"]
sav_r = AO_SUMMARY_ROWS["Total Savings & Debt"]
ao_summary(net_label, lambda l: f"={l}{inc_r}-{l}{exp_r}-{l}{sav_r}")
for col in range(2, 15):
    ao.cell(row=AO_SUMMARY_ROWS[net_label], column=col).fill = fill(GOLD_LT)

print("annual overview ok, summary rows:", AO_SUMMARY_ROWS)

# ---------------------------------------------------------------------------
# DASHBOARD
# ---------------------------------------------------------------------------
dash = wb.create_sheet("Dashboard")
style_sheet_base(dash, tab_color=GOLD)
set_col_widths(dash, {"A": 3, "B": 20, "C": 20, "D": 20, "E": 20, "F": 3, "G": 3})
dash.sheet_view.showGridLines = False

dash.merge_cells("B2:E2")
dash["B2"] = '="YEARLY DASHBOARD — "&' + wb_year_cell
dash["B2"].font = F(size=20, bold=True, color=WHITE)
dash["B2"].fill = fill(TEAL)
dash["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
dash.row_dimensions[2].height = 34

# KPI cards row
kpi_defs = [
    ("TOTAL INCOME (YTD)", f"='Annual Overview'!N{inc_r}", GREEN),
    ("TOTAL EXPENSES (YTD)", f"='Annual Overview'!N{exp_r}", RED),
    ("TOTAL SAVED (YTD)", f"='Annual Overview'!N{sav_r}", TEAL),
    ("NET CASH FLOW (YTD)", f"='Annual Overview'!N{AO_SUMMARY_ROWS[net_label]}", GOLD),
]
col_letters = ["B", "C", "D", "E"]
for (label, formula, color), col in zip(kpi_defs, col_letters):
    dash[f"{col}4"] = label
    dash[f"{col}4"].font = F(size=9, bold=True, color=WHITE)
    dash[f"{col}4"].fill = fill(color)
    dash[f"{col}4"].alignment = CENTER
    dash.row_dimensions[4].height = 22
    dash.merge_cells(f"{col}5:{col}6")
    cell = dash[f"{col}5"]
    cell.value = formula
    cell.number_format = CURRENCY_FMT0
    cell.font = F(size=20, bold=True, color=CHARCOAL)
    cell.alignment = CENTER
    cell.fill = fill(CREAM)
    dash.row_dimensions[5].height = 20
    dash.row_dimensions[6].height = 24

dash["B8"] = "SAVINGS RATE (YTD)"
dash["B8"].font = F(size=9, bold=True, color=WHITE)
dash["B8"].fill = fill(SAGE)
dash.merge_cells("B8:C8")
dash.merge_cells("B9:C10")
dash["B9"] = f"='Annual Overview'!N{sav_r}/'Annual Overview'!N{inc_r}"
dash["B9"].number_format = PCT_FMT
dash["B9"].font = F(size=18, bold=True, color=CHARCOAL)
dash["B9"].alignment = CENTER
dash["B9"].fill = fill(CREAM)

dash.merge_cells("D8:E8")
dash["D8"] = "AVG MONTHLY SPEND"
dash["D8"].font = F(size=9, bold=True, color=WHITE)
dash["D8"].fill = fill(SAGE)
dash.merge_cells("D9:E10")
dash["D9"] = f"='Annual Overview'!N{exp_r}/12"
dash["D9"].number_format = CURRENCY_FMT0
dash["D9"].font = F(size=18, bold=True, color=CHARCOAL)
dash["D9"].alignment = CENTER
dash["D9"].fill = fill(CREAM)

# --- Chart 1: Income vs Expenses vs Savings by month (bar) ---
bar = BarChart()
bar.type = "col"
bar.title = "Income vs. Expenses vs. Savings — By Month"
bar.style = 10
bar.y_axis.title = "Amount ($)"
bar.x_axis.title = "Month"
bar.height = 9
bar.width = 24

cats_ref = Reference(ao, min_col=2, max_col=13, min_row=4, max_row=4)
for label_row, name in ((inc_r, "Income"), (exp_r, "Expenses"), (sav_r, "Savings & Debt")):
    data_ref = Reference(ao, min_col=1, max_col=13, min_row=label_row, max_row=label_row)
    bar.add_data(data_ref, titles_from_data=True, from_rows=True)
bar.set_categories(cats_ref)
bar.legend.position = "b"
dash.add_chart(bar, "B12")

# --- Chart 2: Expense category breakdown (pie), YTD ---
# Build a contiguous helper table (Fixed + Variable categories) off to the side
# so the pie chart has one clean, unbroken data range to point at.
helper_col_label, helper_col_value = 8, 9   # columns H, I
dash.cell(row=11, column=helper_col_label, value="Category").font = F(bold=True)
dash.cell(row=11, column=helper_col_value, value="YTD Total").font = F(bold=True)
helper_start = 12
hr = helper_start
expense_cats_for_pie = FIXED_CATS + VARIABLE_CATS
for cname in expense_cats_for_pie:
    ao_row = AO_CAT_ROWS[cname]
    dash.cell(row=hr, column=helper_col_label, value=cname)
    vcell = dash.cell(row=hr, column=helper_col_value, value=f"='Annual Overview'!N{ao_row}")
    vcell.number_format = CURRENCY_FMT0
    hr += 1
helper_end = hr - 1
for c in (get_column_letter(helper_col_label), get_column_letter(helper_col_value)):
    dash.column_dimensions[c].hidden = True

pie = PieChart()
pie.title = "Where Your Money Went — YTD Expense Breakdown"
pie.height = 9
pie.width = 15
pie_labels = Reference(dash, min_col=helper_col_label, max_col=helper_col_label, min_row=helper_start, max_row=helper_end)
pie_data = Reference(dash, min_col=helper_col_value, max_col=helper_col_value, min_row=11, max_row=helper_end)
pie.add_data(pie_data, titles_from_data=True)
pie.set_categories(pie_labels)
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.dataLabels.showCatName = False
pie.dataLabels.showSerName = False
pie.dataLabels.showVal = False
pie.dataLabels.showLegendKey = False
pie.visible_cells_only = False
dash.add_chart(pie, "K12")

print("dashboard ok")

# ---------------------------------------------------------------------------
# DEBT PAYOFF TRACKER
# ---------------------------------------------------------------------------
debt = wb.create_sheet("Debt Payoff Tracker")
style_sheet_base(debt, tab_color=RED, freeze="A5")
set_col_widths(debt, {"A": 22, "B": 14, "C": 12, "D": 14, "E": 14, "F": 15, "G": 16, "H": 26})

title_block(debt, "DEBT PAYOFF TRACKER", "List every debt, add an extra payment, and see the payoff estimate update instantly.",
            span="A1:H1", color=RED, fontsize=16, height=30)

d_headers = ["Debt Name","Balance","Interest Rate","Minimum Payment","Extra Payment",
             "Total Monthly Payment","Months to Payoff","Notes"]
for i, h in enumerate(d_headers, start=1):
    c = debt.cell(row=3, column=i, value=h)
    c.font = F(bold=True, color=WHITE)
    c.fill = fill(RED)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER_ALL
debt.row_dimensions[3].height = 30

debt_examples = [
    ("Visa Credit Card", 3200, 0.2199, 90, 60),
    ("Student Loan", 14500, 0.055, 165, 0),
    ("Car Loan", 9800, 0.049, 320, 0),
]
DEBT_FIRST_ROW, DEBT_LAST_ROW = 4, 20
r = DEBT_FIRST_ROW
for name, bal, rate, minp, extra in debt_examples:
    debt.cell(row=r, column=1, value=name)
    debt.cell(row=r, column=2, value=bal)
    debt.cell(row=r, column=3, value=rate)
    debt.cell(row=r, column=4, value=minp)
    debt.cell(row=r, column=5, value=extra)
    r += 1

for rr in range(DEBT_FIRST_ROW, DEBT_LAST_ROW + 1):
    for col in (1, 2, 3, 4, 5, 8):
        cell = debt.cell(row=rr, column=col)
        cell.font = F(size=10.5, bold=True, color=INPUT_BLUE)
        cell.fill = fill(GOLD_LT)
        cell.border = BORDER_ALL
    debt.cell(row=rr, column=2).number_format = CURRENCY_FMT
    debt.cell(row=rr, column=3).number_format = PCT_FMT
    debt.cell(row=rr, column=4).number_format = CURRENCY_FMT
    debt.cell(row=rr, column=5).number_format = CURRENCY_FMT
    debt.cell(row=rr, column=1).alignment = LEFT
    debt.cell(row=rr, column=8).alignment = LEFT

    total_pay = debt.cell(row=rr, column=6, value=f"=IF(D{rr}=0,\"\",D{rr}+E{rr})")
    total_pay.number_format = CURRENCY_FMT
    total_pay.font = F(size=10.5, color=CHARCOAL)
    total_pay.border = BORDER_ALL
    total_pay.fill = fill(WHITE)

    months = debt.cell(row=rr, column=7,
        value=(f"=IF(OR(B{rr}=0,F{rr}=0),\"\","
               f"IF(F{rr}<=B{rr}*C{rr}/12,\"Payment too low\","
               f"IF(C{rr}=0,ROUNDUP(B{rr}/F{rr},0),"
               f"ROUNDUP(NPER(C{rr}/12,-F{rr},B{rr}),0))))"))
    months.number_format = '0" months"'
    months.font = F(size=10.5, bold=True, color=RED)
    months.border = BORDER_ALL
    months.fill = fill(WHITE)
    months.alignment = CENTER

debt.cell(row=DEBT_LAST_ROW + 2, column=1, value="TOTAL DEBT").font = F(bold=True, color=CHARCOAL)
tot = debt.cell(row=DEBT_LAST_ROW + 2, column=2, value=f"=SUM(B{DEBT_FIRST_ROW}:B{DEBT_LAST_ROW})")
tot.number_format = CURRENCY_FMT
tot.font = F(bold=True, color=RED, size=12)
debt.cell(row=DEBT_LAST_ROW + 2, column=1).border = Border(top=Side(style="thin", color=CHARCOAL))
tot.border = Border(top=Side(style="thin", color=CHARCOAL))

debt.merge_cells(f"A{DEBT_LAST_ROW+4}:H{DEBT_LAST_ROW+4}")
debt[f"A{DEBT_LAST_ROW+4}"] = ("Tip: pay minimums on everything, then throw every spare dollar at the smallest "
                                "balance (snowball) or the highest interest rate (avalanche) as your Extra Payment.")
debt[f"A{DEBT_LAST_ROW+4}"].font = F(size=9.5, italic=True, color=SAGE)

print("debt tracker ok")

# ---------------------------------------------------------------------------
# SAVINGS GOALS
# ---------------------------------------------------------------------------
sg = wb.create_sheet("Savings Goals")
style_sheet_base(sg, tab_color=SAGE, freeze="A5")
set_col_widths(sg, {"A": 24, "B": 14, "C": 14, "D": 14, "E": 16, "F": 18, "G": 22})

title_block(sg, "SAVINGS GOALS TRACKER", "Turn a wish into a plan — see exactly how much to save each month.",
            span="A1:G1", color=SAGE, fontsize=16, height=30)

s_headers = ["Goal Name","Target Amount","Current Saved","Target Date",
             "% Complete","Monthly Amount Needed","Notes"]
for i, h in enumerate(s_headers, start=1):
    c = sg.cell(row=3, column=i, value=h)
    c.font = F(bold=True, color=WHITE)
    c.fill = fill(SAGE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER_ALL
sg.row_dimensions[3].height = 30

goal_examples = [
    ("Emergency Fund", 10000, 3200, datetime.date(2026, 12, 31)),
    ("Dream Vacation", 4000, 650, datetime.date(2026, 8, 1)),
    ("New Car Down Payment", 6000, 900, datetime.date(2027, 3, 1)),
]
SG_FIRST_ROW, SG_LAST_ROW = 4, 18
r = SG_FIRST_ROW
for name, target, current, tdate in goal_examples:
    sg.cell(row=r, column=1, value=name)
    sg.cell(row=r, column=2, value=target)
    sg.cell(row=r, column=3, value=current)
    sg.cell(row=r, column=4, value=tdate)
    r += 1

today_ref = '"' + datetime.date.today().isoformat() + '"'
for rr in range(SG_FIRST_ROW, SG_LAST_ROW + 1):
    for col in (1, 2, 3, 4, 7):
        cell = sg.cell(row=rr, column=col)
        cell.font = F(size=10.5, bold=True, color=INPUT_BLUE)
        cell.fill = fill(GOLD_LT)
        cell.border = BORDER_ALL
    sg.cell(row=rr, column=2).number_format = CURRENCY_FMT
    sg.cell(row=rr, column=3).number_format = CURRENCY_FMT
    sg.cell(row=rr, column=4).number_format = DATE_FMT
    sg.cell(row=rr, column=1).alignment = LEFT
    sg.cell(row=rr, column=4).alignment = CENTER
    sg.cell(row=rr, column=7).alignment = LEFT

    pct = sg.cell(row=rr, column=5, value=f"=IF(B{rr}=0,\"\",MIN(C{rr}/B{rr},1))")
    pct.number_format = PCT_FMT
    pct.font = F(size=10.5, bold=True, color=TEAL)
    pct.border = BORDER_ALL
    pct.fill = fill(WHITE)
    pct.alignment = CENTER

    monthly = sg.cell(row=rr, column=6,
        value=(f"=IF(OR(B{rr}=0,D{rr}=0),\"\","
               f"IF(D{rr}<=TODAY(),\"Past due\","
               f"IFERROR(MAX((B{rr}-C{rr})/((YEAR(D{rr})-YEAR(TODAY()))*12+(MONTH(D{rr})-MONTH(TODAY()))),0),\"\")))"))
    monthly.number_format = CURRENCY_FMT
    monthly.font = F(size=10.5, bold=True, color=GREEN)
    monthly.border = BORDER_ALL
    monthly.fill = fill(WHITE)

pct_col_range = f"E{SG_FIRST_ROW}:E{SG_LAST_ROW}"
sg.conditional_formatting.add(pct_col_range,
    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                color="2F5D62", showValue=True, minLength=None, maxLength=None))

sg.merge_cells(f"A{SG_LAST_ROW+2}:G{SG_LAST_ROW+2}")
sg[f"A{SG_LAST_ROW+2}"] = "Monthly Amount Needed assumes even monthly contributions from today until your target date."
sg[f"A{SG_LAST_ROW+2}"].font = F(size=9.5, italic=True, color=SAGE)

print("savings goals ok")

# ---------------------------------------------------------------------------
# FINAL POLISH: sheet order, properties, save
# ---------------------------------------------------------------------------
desired_order = (["Welcome", "Dashboard"] + MONTH_ABBR +
                  ["Annual Overview", "Transactions", "Debt Payoff Tracker", "Savings Goals", "Categories"])
wb._sheets = [wb[name] for name in desired_order]
wb.active = 0

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = False

# Trim print areas on long, mostly-blank sheets so a buyer who hits Ctrl+P
# doesn't get a stack of near-empty pages (formulas/validation still work
# past this range - it only limits what prints).
tx.print_area = "A1:H60"
debt.print_area = f"A1:H{DEBT_LAST_ROW + 5}"
sg.print_area = f"A1:G{SG_LAST_ROW + 3}"

wb.properties.title = "The Ultimate Budget Planner"
wb.properties.creator = "Budget Planner"
wb.properties.subject = "Personal Budget Tracker"
wb.properties.description = "Monthly + annual budget planner with dashboard, transaction log, debt payoff and savings goal trackers."

out_path = "Ultimate_Budget_Planner.xlsx"
wb.save(out_path)
print("SAVED", out_path)
